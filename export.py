"""Excel export: order plan, weekly forecasts with ranges, model accuracy, settings used."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from forecast import MODELS, ProductForecast

HEADER = PatternFill("solid", fgColor="1B3A57")
STATUS_FILL = {"Order now": "F8D7D7", "Order soon": "FDEFC8", "OK": "DDF1DD", "Overstocked": "DCE8F7"}


def workbook(plan: pd.DataFrame, forecasts: dict[str, ProductForecast], accuracy: pd.DataFrame,
             settings: pd.DataFrame, notes: list[str]) -> bytes:
    fc_rows = []
    for f in forecasts.values():
        for w, m, lo, hi in zip(f.weeks, f.mean, f.low, f.high):
            fc_rows.append({"Product": f.product, "Week (Monday)": w.date(), "Forecast": round(float(m)),
                            "Low (80 %)": round(float(lo)), "High (80 %)": round(float(hi)), "Model": f.model})
    acc = accuracy.copy()
    for m in MODELS:
        acc[m] = (acc[m] * 100).round(1)
    acc = acc.rename(columns={m: f"{m} error %" for m in MODELS})

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        plan.to_excel(xw, sheet_name="Order plan", index=False, startrow=len(notes) + 1)
        pd.DataFrame(fc_rows).to_excel(xw, sheet_name="Forecast", index=False)
        acc.to_excel(xw, sheet_name="Model accuracy", index=False)
        settings.to_excel(xw, sheet_name="Settings used", index=False)
        ws = xw.sheets["Order plan"]
        for i, n in enumerate(notes, start=1):
            ws.cell(row=i, column=1, value=n).font = Font(italic=i > 1, bold=i == 1)
        for name, sheet in xw.sheets.items():
            hdr = len(notes) + 2 if name == "Order plan" else 1
            for cell in sheet[hdr]:
                cell.fill, cell.font = HEADER, Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for col in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(col)].width = 26 if col == 1 else 15
            sheet.freeze_panes = sheet.cell(row=hdr + 1, column=2)
        status_col = list(plan.columns).index("Status") + 1 if "Status" in plan.columns else None
        pct_cols = [list(plan.columns).index(c) + 1 for c in plan.columns if c.startswith("Served")]
        for r in range(len(notes) + 3, len(notes) + 3 + len(plan)):
            if status_col:
                c = ws.cell(row=r, column=status_col)
                c.fill = PatternFill("solid", fgColor=STATUS_FILL.get(str(c.value), "FFFFFF"))
            for pc in pct_cols:
                ws.cell(row=r, column=pc).number_format = "0%"
    return buf.getvalue()
